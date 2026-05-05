#!/usr/bin/env python3
"""
Generate a client-ready DOCX describing proposed contributions:
- pending items + enhancements
- explicitly includes the 5 requested points

Output:
  docs/MergeTasks_Contribution_Plan.docx
"""

from __future__ import annotations

import os
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "MergeTasks_Contribution_Plan.docx"


def add_title(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(20)


def add_subtitle(doc: Document, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.italic = True
    run.font.size = Pt(11)


def add_h2(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 2")


def add_h3(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="Heading 3")


def add_bullets(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Bullet")


def add_numbered(doc: Document, items: list[str]) -> None:
    for it in items:
        doc.add_paragraph(it, style="List Number")


def main() -> None:
    OUT.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    # Tighter page margins for nicer 1–2 page doc
    for section in doc.sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    add_title(doc, "MergeTasks — Contribution Plan (Enhancements & Pending Work)")
    add_subtitle(
        doc,
        "Purpose: Provide a clear, client-ready list of improvements and remaining work "
        "we can contribute to in the current MergeTasks codebase.",
    )

    add_h2(doc, "Executive Summary")
    doc.add_paragraph(
        "We will focus on improving maintainability, scalability, and operational reliability. "
        "This includes separating frontend and backend into two servers, restructuring parts of the codebase, "
        "productionizing AI as a dedicated service, validating Stripe payments, "
        "and strengthening Docker-based deployment with CI/CD."
    )

    add_h2(doc, "Key Contribution Items (Requested Points)")
    add_numbered(
        doc,
        [
            "Two-server architecture (Frontend + Backend): run the React UI as a dedicated frontend server "
            "(static build/CDN-ready) and the Express+tRPC API as a dedicated backend server with independent "
            "scaling, logging, and health checks.",
            "Unstructured code to structured code: standardize module boundaries, folder conventions, and shared "
            "utilities; reduce cross-module coupling and improve readability and testability.",
            "Build the AI service: implement Copilot as a dedicated AI service using LangGraph, with clear tool "
            "interfaces, memory boundaries, and provider abstraction (Anthropic/Gemini/OpenAI).",
            "Pending: test the payment gateway (requires Stripe API): complete end-to-end validation of Stripe "
            "Connect onboarding + checkout flows in test mode, including webhook verification and failure scenarios.",
            "Easy deployment: use Docker and CI/CD: ensure reproducible builds, environment separation, and automated "
            "deployment pipelines with checks (tests/lints) and health probes.",
        ],
    )

    add_h2(doc, "Enhancements (High Impact)")
    add_h3(doc, "Environment & Configuration")
    add_bullets(
        doc,
        [
            "Split environment files for local/docker/production to prevent misconfiguration (OAuth redirect mismatches, wrong DB).",
            "Add an environment validation/guard script to prevent production secrets from being used in local dev by accident.",
        ],
    )

    add_h3(doc, "Architecture & Modularity")
    add_bullets(
        doc,
        [
            "Define clear domain modules (Clients, Proposals, Stores, Orders, Billing) with strict boundaries.",
            "Introduce shared domain types/contracts to reduce duplication across routers, services, and UI.",
            "Add feature flags for optional subsystems (Stripe, AI providers, email providers) to avoid runtime 500s when not configured.",
        ],
    )

    add_h3(doc, "AI Copilot Reliability & Cost Control")
    add_bullets(
        doc,
        [
            "Add short-TTL caching for AI memory/context building (Redis-backed) to reduce DB load for active sessions.",
            "Improve provider fallback observability (surface which provider/model was used, quota/rate-limit warnings).",
            "Harden tool execution with clearer approval gating and audit trails per organization/user.",
        ],
    )

    add_h3(doc, "Payments (Stripe)")
    add_bullets(
        doc,
        [
            "Add a setup checklist UI for Stripe Connect requirements (platform profile responsibilities, test/live mode alignment).",
            "Implement automated webhook replay verification and local test harness for critical events.",
        ],
    )

    add_h3(doc, "DevOps & Delivery")
    add_bullets(
        doc,
        [
            "Harden Docker compose defaults for local use (ports, env passthrough, health checks).",
            "CI checks: typecheck, tests, lint, and build; publish artifacts and run smoke checks on deploy.",
        ],
    )

    add_h2(doc, "Pending Items (To Close Before Client Handoff)")
    add_bullets(
        doc,
        [
            "Stripe payment gateway testing (requires Stripe API keys + webhooks configured in test mode).",
            "OAuth redirect URI alignment for the chosen local port/domain (Google/Microsoft).",
            "AI provider configuration verification (Anthropic + Gemini) and model availability compatibility.",
            "Email provider verification for non-test sender domains (Resend verified domain) if emails must reach external recipients.",
        ],
    )

    add_h2(doc, "Proposed Delivery Plan (Phased)")
    add_numbered(
        doc,
        [
            "Phase 1 — Stabilize local/dev: environment separation, Docker defaults, basic status/health checks.",
            "Phase 2 — Payments validation: Stripe Connect onboarding + checkout + webhook E2E test plan and execution.",
            "Phase 3 — AI service: LangGraph-based Copilot service, tool contracts, memory boundaries, observability.",
            "Phase 4 — Modularity & microservices path: identify service boundaries, extract the first service (AI or Billing) behind stable APIs.",
        ],
    )

    add_h2(doc, "Notes / Assumptions")
    add_bullets(
        doc,
        [
            "All production secrets should be managed outside the repository (secret manager / CI variables).",
            "Stripe Connect requirements must be accepted in the Stripe Dashboard for onboarding links to succeed.",
            "Microservices migration is recommended as an incremental extraction (strangler pattern), not a rewrite.",
        ],
    )

    doc.add_paragraph()
    doc.add_paragraph("Prepared for: Client Review")
    doc.add_paragraph("Project: MergeTasks")

    doc.save(str(OUT))
    print(str(OUT))


if __name__ == "__main__":
    main()

